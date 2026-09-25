"""Read-only source analysis; every generated file stays in this directory."""
from pathlib import Path
import sys,json,hashlib,collections,platform,importlib.util
import numpy as np
import pandas as pd
from scipy.io import loadmat

OUT=Path(__file__).resolve().parent
ROOT=OUT.parent
SOURCE=ROOT/'2.+Composites'/'2. Composites'/'Layup1'
AUDIT=ROOT/'audit_stiffness'
SPECIMENS=('L1_S11_F','L1_S12_F')

def finite_scalar(v):
    a=np.asarray(v)
    return float(a.item()) if a.size==1 and np.isfinite(a.item()) else np.nan

def survey():
    if (OUT/'acquisition_inventory.csv').exists():
        for r in json.loads((OUT/'source_hashes_before.json').read_text()):
            assert hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()==r['sha256'], 'Source changed: rebuild inventory explicitly in a new output directory.'
        return pd.read_csv(OUT/'acquisition_inventory.csv')
    rows=[];records=[];manifest=[]
    for sp in SPECIMENS:
        for p in sorted((SOURCE/sp/'PZT-data').glob('*.mat')):
            if p.name.startswith('.'):continue
            payload=p.read_bytes()
            manifest.append(dict(path=str(p.relative_to(ROOT)),sha256=hashlib.sha256(payload).hexdigest(),bytes=len(payload)))
            c=loadmat(p,simplify_cells=True)['coupon']
            base=dict(specimen=sp,file=p.name,cycles=finite_scalar(c['cycles']),condition=str(c['condition']).lower().strip(),load_kip=finite_scalar(c['load']),published_index=finite_scalar(c['straingage_data']['stiffness_degradation']))
            records.append(base)
            for index,d in enumerate(c['path_data']):
                a=np.asarray(d['signal_actuator']).ravel();s=np.asarray(d['signal_sensor']).ravel()
                rows.append(dict(**base,array_index=index,actuator=d['actuator'],sensor=d['sensor'],frequency_raw=d['frequency'],gain_raw=d['gain'],amplitude_raw=d['amplitude'],sampling_rate=d['sampling_rate'],actuator_n=len(a),sensor_n=len(s)))
        print('Survey complete:',sp,flush=True)
    frame=pd.DataFrame(rows)
    frame.to_csv(OUT/'acquisition_inventory.csv',index=False)
    pd.DataFrame(records).to_csv(OUT/'record_inventory.csv',index=False)
    (OUT/'source_hashes_before.json').write_text(json.dumps(manifest,indent=2))
    return frame

def describe(frame):
    rec=pd.read_csv(OUT/'record_inventory.csv')
    print('Modules:',{m:bool(importlib.util.find_spec(m)) for m in ['matplotlib','nbformat','nbclient','ipykernel']})
    print('Labelled records by condition:')
    print(rec[rec.published_index.notna()].groupby(['specimen','condition']).agg(recordings=('file','size'),cycles=('cycles','nunique')).to_string())
    print('Settings values:')
    for k in ['frequency_raw','gain_raw','amplitude_raw','sampling_rate','actuator_n','sensor_n']:
        print(k,sorted(frame[k].unique()))
    # Settings and coverage only: no target values used to rank choices.
    summaries=[]
    for condition in ['traction free','clamped','loaded']:
        f=frame[frame.condition.eq(condition)&frame.published_index.notna()]
        cols=['actuator','sensor','frequency_raw','gain_raw','amplitude_raw','sampling_rate','sensor_n','actuator_n']
        for key,g in f.groupby(cols):
            counts=g.groupby('specimen').cycles.nunique()
            if len(counts)==2:
                summaries.append(dict(condition=condition,**dict(zip(cols,key)),L1S11=int(counts.get(SPECIMENS[0],0)),L1S12=int(counts.get(SPECIMENS[1],0))))
    coverage=pd.DataFrame(summaries)
    coverage['min_coverage']=coverage[['L1S11','L1S12']].min(axis=1)
    coverage.to_csv(OUT/'comparable_settings_coverage.csv',index=False)
    print(coverage.sort_values(['min_coverage','actuator','sensor','frequency_raw'],ascending=[False,True,True,True]).head(25).to_string(index=False))
    # Inspect one physically named path across baseline and earliest free repeats.
    for sp in SPECIMENS:
        eligible=rec[(rec.specimen==sp)&rec.condition.isin(['traction free','baseline'])].sort_values(['cycles','file']).head(5)
        for r in eligible.itertuples():
            c=loadmat(SOURCE/sp/'PZT-data'/r.file,simplify_cells=True)['coupon']
            d=next(d for d in c['path_data'] if (d['actuator'],d['sensor'],d['frequency'])==(1,7,150))
            info={'specimen':sp,'file':r.file,'gain':d['gain'],'amplitude':d['amplitude']}
            for channel in ['signal_actuator','signal_sensor']:
                x=np.asarray(d[channel],dtype=float).ravel();xc=x-x.mean()
                peak=np.argmax(np.abs(np.fft.rfft(xc))[1:])+1
                info[channel]=dict(min=float(x.min()),max=float(x.max()),std=float(x.std()),peak_cycles_per_sample=peak/len(x),extreme_fraction=float(max((x==x.min()).mean(),(x==x.max()).mean())),first_samples=x[:12].tolist())
            print(json.dumps(info))

def choose_paths(frame):
    """Use settings + existence of labels, never numerical label magnitudes."""
    f=frame[frame.condition.eq('traction free')&frame.published_index.notna()]
    cols=['actuator','sensor','frequency_raw','gain_raw','amplitude_raw','sampling_rate','sensor_n','actuator_n']
    required=f.groupby('specimen').cycles.nunique().to_dict()
    eligible=[]
    for key,g in f.groupby(cols):
        counts=g.groupby('specimen').cycles.nunique().to_dict()
        if counts==required:
            settings=dict(zip(cols,key))
            # Baseline compatibility is mandatory, including the L1S11 Baseline tag.
            for sp in SPECIMENS:
                name=sp.replace('_','').replace('F','')+'_0_0.mat'
                baseline=frame[(frame.specimen==sp)&(frame.file==name)]
                for k,v in settings.items():baseline=baseline[baseline[k]==v]
                if len(baseline)!=1:break
            else:eligible.append(settings)
    e=pd.DataFrame(eligible)
    counts=e.groupby('frequency_raw').size()
    frequency=sorted(counts.index,key=lambda v:(-counts.loc[v],v))[0]
    selected=e[e.frequency_raw==frequency].sort_values(['actuator','sensor']).head(3)
    plan=dict(condition='traction free; explicitly documented _0_0 baseline also allowed',
        rule='Require all labelled cycles and pristine references in both coupons with exact settings; choose frequency with most eligible paths (tie: lower frequency); take first three actuator/sensor pairs in numeric order.',
        target_values_used_for_selection=False,cycle_max=100000,
        eligible_settings=len(e),eligible_per_frequency={str(k):int(v) for k,v in counts.items()},paths=selected.to_dict('records'),
        reference_rule='Own specimen pristine _0_0 recording; matched actuator, sensor, frequency, gain, amplitude, rate, length; no later reference or best-correlated reference selection.',
        preprocessing='Whole 2000-sample trace, subtract its own mean; no filtering, alignment, resampling, learned normalization or target-selected window.',
        qc=dict(flat_relative_range=1e-10,near_rail_fraction=.01,near_rail_range_tolerance=.001,min_exact_rail_run=3,expected_length=2000))
    (OUT/'selection_plan.json').write_text(json.dumps(plan,indent=2))
    e.to_csv(OUT/'eligible_full_coverage_paths.csv',index=False)
    return plan

def qc_signal(x):
    x=np.asarray(x,dtype=float).ravel(); n=len(x)
    finite=np.isfinite(x).all()
    out=dict(n=n,nonfinite=int((~np.isfinite(x)).sum()),flat=False,clipping_suspect=False)
    if not finite or not n:return out
    lo=x.min();hi=x.max();span=hi-lo
    out['flat']=bool(span<=1e-10*max(1.,abs(lo),abs(hi)))
    def max_run(mask):
        edges=np.diff(np.r_[False,mask,False].astype(int))
        return int(np.max(np.where(edges==-1)[0]-np.where(edges==1)[0],initial=0))
    rail_run=max(max_run(x==lo),max_run(x==hi))
    near_rail=max(float((x<=lo+.001*span).mean()),float((x>=hi-.001*span).mean()))
    out.update(min=float(lo),max=float(hi),std=float(x.std()),rail_run=rail_run,near_rail_fraction=near_rail,
        clipping_suspect=bool(not out['flat'] and (rail_run>=3 or near_rail>=.01)))
    return out

def extract(plan):
    rec=pd.read_csv(OUT/'record_inventory.csv')
    rec=rec[(rec.condition.eq('traction free')|rec.file.str.endswith('_0_0.mat'))&rec.cycles.le(plan['cycle_max'])].copy()
    signal_rows=[];features=[];cache={};spectra=[]
    for sp in SPECIMENS:
        basefile=sp.replace('_','').replace('F','')+'_0_0.mat'
        records=rec[rec.specimen==sp].sort_values(['cycles','file'])
        references={}
        for r in records.itertuples():
            c=loadmat(SOURCE/sp/'PZT-data'/r.file,simplify_cells=True)['coupon']
            lookup=collections.defaultdict(list)
            for d in c['path_data']:lookup[(d['actuator'],d['sensor'],d['frequency'])].append(d)
            for setting in plan['paths']:
                key=(setting['actuator'],setting['sensor'],setting['frequency_raw'])
                assert len(lookup[key])==1,('Missing/duplicate path',r.file,key)
                d=lookup[key][0];path=f"A{key[0]}-S{key[1]}"
                a=np.asarray(d['signal_actuator'],dtype=float).ravel();s=np.asarray(d['signal_sensor'],dtype=float).ravel()
                match=all(d[k]==setting[k+'_raw'] for k in ['gain','amplitude']) and d['sampling_rate']==setting['sampling_rate'] and len(a)==setting['actuator_n'] and len(s)==setting['sensor_n']
                qs=qc_signal(s);qa=qc_signal(a)
                base=dict(specimen=sp,file=r.file,cycles=r.cycles,condition=r.condition,load_kip=r.load_kip,published_index=r.published_index,path=path,actuator=key[0],sensor=key[1],frequency_raw=key[2],gain_raw=d['gain'],amplitude_raw=d['amplitude'],sampling_rate=d['sampling_rate'],settings_match=match)
                for channel,q in [('sensor',qs),('actuator',qa)]:signal_rows.append(dict(**base,channel=channel,**q))
                x=s-s.mean();ax=a-a.mean()
                peak=np.argmax(np.abs(np.fft.rfft(ax))[1:])+1
                spectra.append(dict(**base,peak_bin=int(peak),peak_cycles_per_sample=peak/len(a),actuator_fft_peak_hz_metadata=peak*d['sampling_rate']/len(a),nominal_excitation_hz=key[2]*1000,inferred_rate_if_nominal_tone_is_correct=key[2]*1000/(peak/len(a))))
                if r.file==basefile:references[path]=(x,ax,setting)
                ref,aref,_=references[path]
                rms=float(np.sqrt(np.mean(x*x))); rms_ref=float(np.sqrt(np.mean(ref*ref)))
                energy=float(np.sum(x*x));energy_ref=float(np.sum(ref*ref))
                corr=float(np.dot(x,ref)/(np.linalg.norm(x)*np.linalg.norm(ref))) if np.linalg.norm(x)>0 else np.nan
                # Rectangle-rule integral over 2000 sample bins: conditional on exported rate.
                row=dict(**base,reference_file=basefile,is_reference=r.file==basefile,
                    sensor_rms_au=rms,sensor_sum_squares_au2=energy,
                    sensor_integral_metadata_au2_s=energy/d['sampling_rate'],
                    rms_ratio=rms/rms_ref,energy_ratio=energy/energy_ref,reference_correlation=corr,
                    actuator_rms_ratio=float(np.linalg.norm(ax)/np.linalg.norm(aref)),
                    valid_basic=bool(match and not qs['nonfinite'] and not qa['nonfinite'] and not qs['flat'] and not qa['flat']),
                    clipping_suspect=bool(qs['clipping_suspect'] or qa['clipping_suspect']),
                    sensor_clipping_suspect=qs['clipping_suspect'],actuator_clipping_suspect=qa['clipping_suspect'])
                features.append(row)
                cache[sp+'__'+r.file+'__'+path+'__sensor']=s
                cache[sp+'__'+r.file+'__'+path+'__actuator']=a
        print('Features extracted:',sp,flush=True)
    pd.DataFrame(signal_rows).to_csv(OUT/'signal_qc.csv',index=False)
    f=pd.DataFrame(features)
    f['filename_repeat_suffix']=f.file.str.extract(r'_\d+_(\d+)\.mat')[0]
    f['repeat_id']=np.where(f.is_reference,'baseline',f.filename_repeat_suffix)
    f.to_csv(OUT/'features.csv',index=False)
    pd.DataFrame(spectra).to_csv(OUT/'sampling_diagnostics.csv',index=False)
    np.savez_compressed(OUT/'selected_signals.npz',**cache)
    print('FEATURE SUMMARY',f.groupby('specimen').agg(records=('file','nunique'),checkpoints=('cycles','nunique'),feature_rows=('file','size'),setting_mismatches=('settings_match',lambda x:int((~x).sum())),basic_failures=('valid_basic',lambda x:int((~x).sum())),clipping_flags=('clipping_suspect','sum')).to_string())
    print('Clipping flags:',f[f.clipping_suspect][['specimen','file','path','sensor_clipping_suspect','actuator_clipping_suspect']].to_string(index=False))
    return f

def enrich_and_summarize(f,frame):
    from openpyxl import load_workbook
    from datetime import datetime,time
    from scipy.stats import spearmanr
    logrows=[]
    for sp in SPECIMENS:
        book=next(p for p in (SOURCE/sp).glob('*.xlsx') if not p.name.startswith('.'))
        ws=load_workbook(book,data_only=True)['Log_fatigue_testing']
        for row in ws.iter_rows(min_row=2,values_only=False):
            v=[c.value for c in row]
            if len(v)<9 or not isinstance(v[6],str):continue
            name=v[6].strip(" '\"")
            when=None
            if isinstance(v[0],datetime) and isinstance(v[1],time):when=datetime.combine(v[0].date(),v[1]).isoformat(sep=' ')
            logrows.append(dict(specimen=sp,file=name,logbook=book.name,logbook_row=row[0].row,acquired_at=when,log_date=str(v[0]),log_time=str(v[1]),log_cycle=v[3],log_condition=v[5],mts_file=v[7],remark=v[8]))
    logs=pd.DataFrame(logrows)
    assert not logs.duplicated(['specimen','file']).any()
    f=f.merge(logs,on=['specimen','file'],how='left',validate='many_to_one')
    assert f.logbook_row.notna().all()
    assert f.cycles.eq(pd.to_numeric(f.log_cycle)).all()
    # Filename condition codes are not authoritative (e.g. L1S11_1000_1_2).
    # Assign unique display repeat ordinals in worksheet order within each cycle.
    recordings=f.drop_duplicates(['specimen','file']).sort_values(['specimen','cycles','logbook_row','file'])
    nonref=recordings[~recordings.is_reference].copy()
    nonref['repeat_ordinal']=(nonref.groupby(['specimen','cycles']).cumcount()+1).astype(str)
    ordinal=dict(zip(zip(nonref.specimen,nonref.file),nonref.repeat_ordinal))
    f['repeat_id']=['baseline' if r.is_reference else ordinal[(r.specimen,r.file)] for r in f.itertuples()]
    f.to_csv(OUT/'features.csv',index=False)
    logs[(logs.specimen==SPECIMENS[1])&pd.to_numeric(logs.log_cycle,errors='coerce').between(10000,30000)].to_csv(OUT/'L1S12_jump_logbook_rows.csv',index=False)
    metrics=['rms_ratio','energy_ratio','reference_correlation']
    labelled=f[f.published_index.notna()&f.valid_basic&~f.clipping_suspect]
    coverage=labelled.groupby('specimen').agg(labelled_recordings=('file','nunique'),labelled_checkpoints=('cycles','nunique'),feature_rows=('file','size'))
    coverage.to_csv(OUT/'usable_coverage.csv')
    g=f.groupby(['specimen','cycles','path'])
    med=g[metrics+['published_index']].median().reset_index()
    med.to_csv(OUT/'checkpoint_medians.csv',index=False)
    ranges=g[metrics].agg(lambda x:float(x.max()-x.min())).reset_index()
    ranges=ranges.merge(g.file.nunique().rename('repeat_count').reset_index(),on=['specimen','cycles','path'])
    ranges.to_csv(OUT/'repeat_ranges.csv',index=False)
    association=[]
    for (sp,path),group in med[med.published_index.notna()].groupby(['specimen','path']):
        for metric in metrics:
            association.append(dict(specimen=sp,path=path,feature=metric,n_checkpoints=len(group),spearman_index=float(spearmanr(group[metric],group.published_index).statistic),spearman_cycles=float(spearmanr(group[metric],group.cycles).statistic)))
    pd.DataFrame(association).to_csv(OUT/'descriptive_associations.csv',index=False)
    # All-path metadata comparison around the gap, not merely the chosen paths.
    names=['L1S12_10000_3_1.mat','L1S12_10000_3_2.mat','L1S12_20000_3_1.mat','L1S12_20000_3_2.mat','L1S12_20000_3_3.mat','L1S12_30000_3_1.mat','L1S12_30000_3_2.mat']
    settings=['gain_raw','amplitude_raw','sampling_rate','sensor_n','actuator_n']
    keys=['actuator','sensor','frequency_raw']
    anchor=frame[(frame.specimen==SPECIMENS[1])&(frame.file==names[0])].set_index(keys)
    changes=[]
    for name in names:
        current=frame[(frame.specimen==SPECIMENS[1])&(frame.file==name)].set_index(keys)
        assert anchor.index.equals(current.index)
        changed=(anchor[settings]!=current[settings]).any(axis=1)
        changes.append(dict(file=name,relative_to=names[0],paths_checked=len(current),changed_settings_paths=int(changed.sum()),**{s+'_changes':int((anchor[s]!=current[s]).sum()) for s in settings}))
    pd.DataFrame(changes).to_csv(OUT/'L1S12_jump_setting_changes.csv',index=False)
    f[(f.specimen==SPECIMENS[1])&f.cycles.between(10000,30000)].to_csv(OUT/'L1S12_jump_features.csv',index=False)
    # Quantify array-order variation without using array positions for any match.
    order=frame.groupby(keys).array_index.nunique()
    print('Paths with varying array position across files:',int((order>1).sum()))
    print('Usable coverage:\n'+coverage.to_string())
    print('Associations (descriptive only):\n'+pd.DataFrame(association).to_string(index=False))
    print('Jump metadata comparisons:\n'+pd.DataFrame(changes).to_string(index=False))
    print('Actuator RMS ratio range:',f.actuator_rms_ratio.min(),f.actuator_rms_ratio.max())
    return f

def sampling_check():
    records=[]
    for sp in SPECIMENS:
        name=sp.replace('_','').replace('F','')+'_0_0.mat'
        c=loadmat(SOURCE/sp/'PZT-data'/name,simplify_cells=True)['coupon']
        for d in c['path_data']:
            if (d['actuator'],d['sensor'])!=(1,10):continue
            for chan in ['signal_actuator','signal_sensor']:
                x=np.asarray(d[chan],dtype=float).ravel();x-=x.mean()
                peak=int(np.argmax(np.abs(np.fft.rfft(x))[1:])+1)
                # Whole-record FFT is a diagnostic, not an exact tone-frequency estimator.
                records.append(dict(specimen=sp,file=name,channel=chan,nominal_khz=d['frequency'],peak_cycles_per_sample=peak/len(x),peak_khz_at_metadata_rate=peak*d['sampling_rate']/len(x)/1000,metadata_rate_hz=d['sampling_rate']))
    pd.DataFrame(records).to_csv(OUT/'baseline_frequency_check.csv',index=False)

def plots(f,plan):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FixedLocator,FuncFormatter
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.titlesize':11,'axes.labelsize':10,'axes.spines.top':False,'axes.spines.right':False,'savefig.facecolor':'white'})
    colors={'A1-S10':'#2166AC','A2-S8':'#B66B17','A2-S11':'#697842'}
    markers={'baseline':'*','1':'o','2':'s','3':'^'}
    metric_labels={'rms_ratio':'RMS / pristine-reference RMS','energy_ratio':'Integrated squared amplitude / reference','reference_correlation':'Correlation with pristine reference'}
    for sp in SPECIMENS:
        sf=f[f.specimen.eq(sp)]
        fig,axes=plt.subplots(3,2,figsize=(13,10),layout='constrained')
        for i,(metric,title) in enumerate(metric_labels.items()):
            for path,pf in sf.groupby('path'):
                for repeat,rf in pf.groupby('repeat_id'):
                    marker=markers[str(repeat)]
                    labelled=rf[rf.published_index.notna()]
                    missing=rf[rf.published_index.isna()]
                    axes[i,0].scatter(labelled.cycles,labelled[metric],s=37,c=colors[path],marker=marker,alpha=.85)
                    if len(missing):axes[i,0].scatter(missing.cycles,missing[metric],s=42,facecolors='none',edgecolors=colors[path],marker=marker,linewidths=1.2)
                    axes[i,1].scatter(labelled.published_index,labelled[metric],s=37,c=colors[path],marker=marker,alpha=.85)
            axes[i,0].set_xscale('symlog',linthresh=1,linscale=.45)
            ticks=[0,1,10,100,1000,10000,100000]
            axes[i,0].set_xticks(ticks,labels=['0','1','10','100','1k','10k','100k'])
            axes[i,0].set_xlabel('Fatigue cycles (symmetric log scale; no interpolation)')
            axes[i,1].set_xlabel('Published stiffness-degradation index')
            axes[i,1].xaxis.set_major_formatter(FuncFormatter(lambda v,p:f'{v:.2f}'))
            for ax in axes[i]:
                ax.set_ylabel(title);ax.grid(alpha=.18);ax.set_axisbelow(True)
                if metric=='reference_correlation':ax.set_ylim(-1.05,1.05)
                else:ax.set_ylim(bottom=0)
            # Comparable scales within each row and across specimens.
            if metric!='reference_correlation':
                ymax=float(f[metric].max())*1.08
                for ax in axes[i]:ax.set_ylim(0,ymax)
        handles=[plt.Line2D([],[],color=colors[p],marker='o',linestyle='',label=p) for p in colors]
        handles += [plt.Line2D([],[],color='#444444',marker=m,linestyle='',label=('Pristine reference' if r=='baseline' else 'Repeat '+r)) for r,m in markers.items()]
        fig.legend(handles=handles,loc='upper center',bbox_to_anchor=(.5,.94),ncol=7,frameon=False)
        nlab=sf[sf.published_index.notna()].cycles.nunique()
        fig.suptitle(f'{sp.replace("_", "").replace("F", "")} | traction free | nominal 450 kHz\n{sf.file.nunique()} recordings; {nlab} labelled checkpoints; own-specimen pristine reference',fontsize=14,y=.995)
        fig.get_layout_engine().set(rect=(0,.045,1,.875))
        fig.text(.02,.015,'Open markers: measured PZT with no label (cycle view only). Shapes retain recording repeats; no fitted lines. Energy ratio = RMS ratio squared.',fontsize=9)
        stem=sp.replace('_','').replace('F','')+'_features'
        fig.savefig(OUT/(stem+'.png'),dpi=160);fig.savefig(OUT/(stem+'.svg'));plt.close(fig)
    # Raw waveform inspection in sample coordinates avoids asserting a physical clock.
    signals=np.load(OUT/'selected_signals.npz')
    fig,axes=plt.subplots(3,2,figsize=(13,8),layout='constrained')
    for j,sp in enumerate(SPECIMENS):
        short=sp.replace('_','').replace('F','')
        for i,path in enumerate(colors):
            ax=axes[i,j]
            for file,color,style,label in [(short+'_0_0.mat','#777777','-','Pristine'),(short+'_100000_3_1.mat','#2166AC','-','100k repeat 1')]:
                key=sp+'__'+file+'__'+path+'__sensor'
                if key not in signals:continue
                x=signals[key];ax.plot(np.arange(len(x)),x-x.mean(),color=color,ls=style,lw=.7,label=label,alpha=.85)
            ax.set_title(short+' '+path);ax.set_xlabel('Sample index');ax.set_ylabel('Centered sensor amplitude (a.u.)');ax.grid(alpha=.15)
    axes[0,0].legend(frameon=False,ncol=2)
    fig.suptitle('Selected sensor waveforms | matching acquisition settings',fontsize=14)
    fig.savefig(OUT/'waveform_check.png',dpi=160);plt.close(fig)
    # First pristine actuator and its spectrum, independent of fatigue target values.
    fig,axes=plt.subplots(1,2,figsize=(12,4),layout='constrained')
    a=signals['L1_S11_F__L1S11_0_0.mat__A1-S10__actuator'];a=a-a.mean()
    axes[0].plot(np.arange(len(a)),a,color='#2166AC',lw=.9)
    axes[0].set(xlim=(0,700),xlabel='Sample index',ylabel='Centered actuator amplitude (a.u.)',title='Nominal 450 kHz excitation')
    freq=pd.read_csv(OUT/'baseline_frequency_check.csv')
    for (sp,channel),g in freq.groupby(['specimen','channel']):
        axes[1].plot(g.nominal_khz,g.peak_khz_at_metadata_rate,color=('#2166AC' if sp==SPECIMENS[0] else '#B66B17'),marker=('o' if channel=='signal_actuator' else 's'),ls=('solid' if channel=='signal_actuator' else 'dashed'),label=sp.replace('_','').replace('F','')+' '+channel.replace('signal_',''),lw=1)
    axes[1].plot([150,450],[150,450],color='#888888',ls=':',label='Equality to nominal')
    axes[1].set(xlabel='Nominal excitation frequency (kHz)',ylabel='FFT peak using exported 1.2 MHz rate (kHz)',title='Timing metadata fails the frequency check')
    axes[1].legend(frameon=False,fontsize=8,ncol=2)
    for ax in axes:ax.grid(alpha=.15)
    fig.savefig(OUT/'sampling_check.png',dpi=160);plt.close(fig)
    # Jump diagnostics: explicit recording IDs, including unlabelled 20k measurements.
    sf=f[(f.specimen==SPECIMENS[1])&f.cycles.between(10000,30000)]
    records=sf.drop_duplicates('file').sort_values(['cycles','file'])
    fig,axes=plt.subplots(3,1,figsize=(12,9),sharex=True,layout='constrained')
    xpos={v:i for i,v in enumerate(records.file)}
    for ax,(metric,title) in zip(axes,metric_labels.items()):
        for path,g in sf.groupby('path'):
            ax.scatter(g.file.map(xpos),g[metric],color=colors[path],label=path,s=45)
        ax.set_ylabel(title);ax.grid(alpha=.18)
        if metric=='reference_correlation':ax.set_ylim(-1.05,1.05)
    labels=[f'{int(r.cycles/1000)}k / repeat {r.repeat_id}\n{str(r.acquired_at)[:10]}' for r in records.itertuples()]
    axes[-1].set_xticks(range(len(labels)),labels);axes[-1].set_xlabel('Recordings in cycle/repeat order (not a uniform time axis)')
    axes[0].legend(frameon=False,ncol=3)
    fig.suptitle('L1S12 | measured features across the 10k–30k interval\n20k index missing; all measurements and the published jump retained',fontsize=14)
    fig.savefig(OUT/'L1S12_jump.png',dpi=160);plt.close(fig)

def verify(f):
    assert len(f)==189 and f.file.nunique()==63
    assert f.load_kip.eq(0).all()
    assert f.repeat_id.eq('baseline').eq(f.is_reference).all()
    recordings=f.drop_duplicates(['specimen','file'])
    assert not recordings[~recordings.is_reference].duplicated(['specimen','cycles','repeat_id']).any()
    assert f.valid_basic.all() and not f.clipping_suspect.any()
    assert np.allclose(f.energy_ratio,f.rms_ratio**2,rtol=1e-12,atol=1e-12)
    assert np.allclose(f.sensor_sum_squares_au2,2000*f.sensor_rms_au**2,rtol=1e-12)
    assert np.allclose(f[f.is_reference].reference_correlation,1)
    assert f.reference_correlation.between(-1.000000001,1.000000001).all()
    assert f[f.published_index.notna()].groupby('specimen').cycles.nunique().to_dict()=={SPECIMENS[0]:15,SPECIMENS[1]:14}
    assert f[(f.specimen==SPECIMENS[1])&(f.cycles==20000)].published_index.isna().all()
    manifests=json.loads((OUT/'source_hashes_before.json').read_text())
    changed=[r['path'] for r in manifests if hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()!=r['sha256']]
    assert not changed
    old=json.loads((AUDIT/'original_manifest.json').read_text())
    dataset=SOURCE.parent
    metadata_changed=[r['path'] for r in old if not (dataset/r['path']).exists() or (dataset/r['path']).stat().st_size!=r['size'] or (dataset/r['path']).stat().st_mtime_ns!=r['mtime_ns']]
    assert not metadata_changed
    result=dict(pzt_sha256_files_checked=len(manifests),pzt_hash_changes=changed,original_metadata_files_checked=len(old),original_metadata_changes=metadata_changed,feature_identity_checks='passed',missing_labels_preserved=True,models_trained=0)
    (OUT/'verification.json').write_text(json.dumps(result,indent=2))
    import scipy,matplotlib,openpyxl
    (OUT/'environment.json').write_text(json.dumps(dict(python=platform.python_version(),numpy=np.__version__,pandas=pd.__version__,scipy=scipy.__version__,matplotlib=matplotlib.__version__,openpyxl=openpyxl.__version__),indent=2))
    print('Verification passed:',json.dumps(result))

def build_notebook():
    import nbformat
    from nbclient import NotebookClient
    md=nbformat.v4.new_markdown_cell;code=nbformat.v4.new_code_cell
    nb=nbformat.v4.new_notebook(cells=[
        md('# L1S11 / L1S12: reproducible PZT feature feasibility\n\n'
           '**Assessment:** simple regression is worth a limited trial. All 29 published checkpoints remain in the matched three-path subset, but there are only two independent coupons. RMS has strong within-specimen associations; waveform correlation is repeat-sensitive. The L1S12 jump is retained and the physical sample rate remains unresolved.\n\n'
           'No interpolation, trained model, target-ranked path selection, or modification of original files. See `ASSESSMENT.md` for full evidence and limits.'),
        md('## Context and methods\n\n'
           'Use traction-free recordings and the pristine zero-load baseline. Survey acquisition metadata before selecting paths. Require full labelled-cycle coverage and exact baseline compatibility, select the frequency with most eligible paths, then the first three actuator/sensor pairs numerically. Numerical target values play no part in selection.\n\n'
           '**Assumptions:** each specimen has its own initial calibration waveform available; same-path matching settings and sample grids support sample-domain comparison. We do not assume the exported clock is physically correct. Missing targets stay missing.'),
        code("from pathlib import Path\nimport sys, io, contextlib, json\nimport pandas as pd\nfrom IPython.display import display, Image\nout = Path.cwd() if (Path.cwd()/'run_feasibility.py').exists() else Path.cwd()/'pzt_feasibility'\nsys.path.insert(0, str(out.resolve()))\nimport run_feasibility as pzt\n"),
        md('## 1. Survey and select by acquisition metadata\n\nThe inventory reads all 254 PZT files on first run. Later runs verify every source SHA-256 before reusing it. All downstream signals are read again from the original MAT files. Paths are keyed by actuator, sensor and frequency, never by array index.'),
        code("inventory = pzt.survey()\nplan = pzt.choose_paths(inventory)\ndisplay(pd.DataFrame(plan['paths']))\nprint('Eligible settings by nominal frequency:', plan['eligible_per_frequency'])\nprint('Selection rule:', plan['rule'])\n"),
        md('## 2. Extract signals, screen quality and calculate features\n\nEach whole trace is centered by its own mean. RMS = sqrt(mean(x²)); the conditional rectangle-rule integral = sum(x²)/exported_fs; zero-lag correlation = dot(x,r)/(norm(x) norm(r)). The primary amplitude features divide by the same-path pristine reference. No filtering, time alignment or target-selected window is used.\n\nClipping checks are heuristics because hardware rails are unknown. See `selection_plan.json` for fixed thresholds.'),
        code("with contextlib.redirect_stdout(io.StringIO()):\n    features = pzt.extract(plan)\n    features = pzt.enrich_and_summarize(features, inventory)\n    pzt.sampling_check()\ndisplay(pd.read_csv(out/'usable_coverage.csv'))\nqc = pd.read_csv(out/'signal_qc.csv')\ndisplay(qc.groupby('channel').agg(traces=('file','size'),nonfinite_samples=('nonfinite','sum'),flat_traces=('flat','sum'),clipping_suspicions=('clipping_suspect','sum')))\nprint('Lengths:', sorted(qc.n.unique()))\nprint('Setting mismatches:', int((~features.settings_match).sum()))\n"),
        md('## 3. Verify sampling units and reference compatibility\n\nNominal excitation uses kHz according to the README; the export hard-codes 1.2 MHz. At nominal 450 kHz the pristine excitation FFT peak is 46.2–46.8 kHz if interpreted using that rate. Across all seven nominal frequencies there is a similar roughly tenfold discrepancy. A near-12-MHz clock is an inference, not an adopted correction. Absolute time integrals remain conditional; reference ratios cancel the shared clock scale. Signal amplitudes are treated as arbitrary exported units.\n\nReferences are each specimen\'s pristine `_0_0.mat`, not averages of later cycle-zero static-loading records. Their path, settings, length and rate match.'),
        code("pzt.plots(features, plan)\ndisplay(Image(filename=str(out/'sampling_check.png'), width=1000))\n"),
        md('## 4. Features by specimen, cycle and published index\n\nAll repetitions remain individual points. Marker shape identifies repeat number and CSV filenames retain exact identities. Open points have a measured PZT feature but no published target; they appear only on the cycle axis. No curves are fitted or interpolated. Amplitude and correlation scales match across specimen figures.'),
        code("display(Image(filename=str(out/'L1S11_features.png'), width=1000))\n"),
        code("display(Image(filename=str(out/'L1S12_features.png'), width=1000))\n"),
        md('RMS decreases within each coupon. Integrated squared-amplitude ratio is exactly RMS ratio squared per recording, so these are redundant predictors. Signed zero-lag correlation is less stable and must not be interpreted as a monotonic damage measure.'),
        md('## 5. L1S12 10k–30k interval: facts and hypotheses\n\n'
           'The index rises from 0.084112497 to 0.233175942; 20k is unlabelled. Logbook rows 53–71 document a September–November gap, temporary absence of strain gauges, later rosette recordings and observed vertical delamination growth. All stored gain, drive-amplitude, sample-rate and length fields remain unchanged across all 252 paths in the seven traction-free records.\n\n'
           'At 30k, the operator notes a signal shift and suggests dye or temperature. Those causes remain hypotheses. Damage growth, gauge/calibration changes and environmental effects cannot be separated here. No automatic deletion or smoothing is performed.'),
        code("display(pd.read_csv(out/'L1S12_jump_setting_changes.csv')[['file','paths_checked','changed_settings_paths']])\ndisplay(Image(filename=str(out/'L1S12_jump.png'), width=1000))\n"),
        md('At A2→S11, 30k-repeat reference correlations are +0.37745 and -0.66720 while RMS ratios are 0.33068 and 0.33338. A shared stiffness target therefore coexists with markedly different phase-sensitive similarity. The unlabelled 20k repeats show variability without an assigned target.'),
        md('## 6. Waveform inspection and descriptive associations\n\nThe waveforms below compare matched pristine and 100k acquisitions in sample coordinates. Associations use one median per checkpoint/path, not thousands of path samples as independent observations. These statistics are descriptive; they were calculated only after selection and never used to choose paths.'),
        code("display(Image(filename=str(out/'waveform_check.png'), width=1000))\nassoc = pd.read_csv(out/'descriptive_associations.csv')\ndisplay(assoc[assoc.feature.ne('energy_ratio')].round(3))\n"),
        md('## 7. Reproducibility and preservation checks\n\nChecks cover label coverage, exact reference identities, amplitude-feature algebra, target missingness, zero recorded load, correlation bounds and source preservation. SHA-256 compares every surveyed PZT file; the earlier inventory also checks all 2,571 original file sizes/timestamps. Package versions are saved in `environment.json`.'),
        code("pzt.verify(features)\ndisplay(pd.Series(json.loads((out/'environment.json').read_text()), name='Version'))\n"),
        md('## Takeaways\n\n'
           '- 15 L1S11 and 14 L1S12 labelled checkpoints remain; 58 labelled recordings are repeated views, not 58 independent specimens.\n'
           '- No selected trace failed the basic screens, but physical timing calibration and absolute amplitude calibration are unresolved.\n'
           '- Try a predeclared low-dimensional regularized regression with RMS ratios **or** energy ratios, compared with a cycle-only baseline. Treat correlation as a separate sensitivity analysis.\n'
           '- Hold out an entire coupon and report both transfer directions; never select the best path from held-out target performance. For forecasting, add forward-time splits.\n'
           '- Preserve the L1S12 jump, the missing 20k label and identifiable repeats. Results cannot establish broad generalization from two coupons.\n\n'
           'Controlling sources: local Layup1 README, Technical Report export code, the two logbooks, original MAT fields and the existing audit. No new labels were manufactured.'),
    ],metadata={'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'}})
    path=OUT/'feasibility.ipynb'
    nbformat.validate(nb)
    nbformat.write(nb,path)
    client=NotebookClient(nb,timeout=180,kernel_name='python3',resources={'metadata':{'path':str(ROOT)}})
    client.execute()
    nbformat.validate(nb);nbformat.write(nb,path)
    errors=[o for cell in nb.cells if cell.cell_type=='code' for o in cell.get('outputs',[]) if o.output_type=='error']
    assert not errors
    print('Executed notebook saved:',path,flush=True)

if __name__=='__main__':
    build_notebook()
